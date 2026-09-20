#!/usr/bin/env python3
"""Парсинг сборников ГЭСН (PDF ФГИС ЦС) в иерархический классификатор работ.

Иерархия класса:
  вид норм (ГЭСН / ГЭСНм / ГЭСНр / ГЭСНп / ГЭСНмр)
    → сборник
      → отдел / раздел / подраздел
        → таблица (группа работ)
          → норма (конкретная работа)

Пример:
  python scripts/parse_gesn_classifier.py
  python scripts/parse_gesn_classifier.py --max-files 3 --output /tmp/gesn-test
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Iterator

import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = ROOT / "dataset" / "documents" / "Нормативная база для классификатора работ"
DEFAULT_OUTPUT = ROOT / "dataset" / "classifier"

DOMAIN_BY_FOLDER = {
    "строительные работы и специальные строительные работы": {
        "id": "gesn",
        "kind": "ГЭСН",
        "name": "Строительные и специальные строительные работы",
    },
    "монтаж оборудования": {
        "id": "gesnm",
        "kind": "ГЭСНм",
        "name": "Монтаж оборудования",
    },
    "ремонтно-строительные работы": {
        "id": "gesnr",
        "kind": "ГЭСНр",
        "name": "Ремонтно-строительные работы",
    },
    "пусконаладочные работы": {
        "id": "gesnp",
        "kind": "ГЭСНп",
        "name": "Пусконаладочные работы",
    },
    "капитальный ремонт оборудования": {
        "id": "gesnmr",
        "kind": "ГЭСНмр",
        "name": "Капитальный ремонт оборудования",
    },
}

FILENAME_RE = re.compile(
    r"Сборник\s+(ГЭСН[мрп]*)(\d+)\s+(?P<name>.+)\.pdf$",
    re.IGNORECASE,
)
DOC_CODE_RE = re.compile(r"(ГЭСН[мрп]*)\s+(81-\d{2}-\d{2}-\d{4})")
COLLECTION_RE = re.compile(r"Сборник\s+(\d+)\.\s*(.+)")
HEADER_LINE_RE = re.compile(r"^Сведения сформированы ФГИС ЦС.*$", re.MULTILINE)
RUNNING_HEADER_RE = re.compile(r"^ГЭСН[мрп]*\s+81-\d{2}-\d{2}-\d{4}\s+\S+.+$", re.MULTILINE)
PAGE_NUMBER_RE = re.compile(r"^\d{1,4}$")
SOFT_HYPHEN_RE = re.compile(r"[\u00ad\u200b\x02]")
TABLE_RE = re.compile(r"Таблица\s+ГЭСН([мрп]*)\s*")
TABLE_CODE_RE = re.compile(r"\d{2}-\d{2}-\d{3}")
RATE_CODE_RE = re.compile(r"^(\d{2}-\d{2}-\d{3}-\d{2})\s+(.*)$")
RATE_WRAP_RE = re.compile(r"(\d{2}-\d{2}-)\s*\n\s*(\d{3}-\d{2})")
TITLE_WORD_RE = re.compile(r'^[«"„(]*[А-ЯЁ][А-Яа-яЁё\-]{3,}')
OTDEL_RE = re.compile(r"^Отдел\s+(\d+)\.\s*(.+)$")
RAZDEL_RE = re.compile(r"^Раздел\s+(\d+)\.\s*(.+)$")
PODRAZDEL_RE = re.compile(r"^Подраздел\s+([\d.]+)\.\s*(.+)$")
APPENDIX_RE = re.compile(r"\n(?:IV\.\s*ПРИЛОЖ|ПРИЛОЖЕНИЯ\s*\n)")
MARKER_LINE_RE = re.compile(
    r"^(?:Отдел\s+\d+\.|Раздел\s+\d+\.|Подраздел\s+[\d.]+\.|Таблица\s+ГЭСН|"
    r"I{1,3}\.\s|Состав работ|Измеритель:|Код ресурса|ПРИЛОЖЕН)"
)
COMPOSITION_ITEM_RE = re.compile(r"^(\d{2})\.\s+(.*)$")


def normalize_ws(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def domain_from_folder(folder_name: str) -> dict[str, str]:
    lowered = folder_name.lower()
    for key, meta in DOMAIN_BY_FOLDER.items():
        if key in lowered:
            return dict(meta)
    return {"id": "other", "kind": "ГЭСН", "name": folder_name}


def is_mostly_upper(text: str) -> bool:
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return False
    upper = sum(1 for c in letters if c.isupper())
    return upper / len(letters) >= 0.7


def extract_pdf_text(path: Path) -> str:
    pdf = pdfium.PdfDocument(str(path))
    pages: list[str] = []
    document_codes: list[str] = []
    try:
        for i in range(len(pdf)):
            raw = pdf[i].get_textpage().get_text_bounded()
            raw = raw.replace("\r\n", "\n").replace("\r", "\n")
            raw = SOFT_HYPHEN_RE.sub("", raw)
            for match in DOC_CODE_RE.finditer(raw):
                document_codes.append(f"{match.group(1)} {match.group(2)}")
            raw = HEADER_LINE_RE.sub("", raw)
            raw = RUNNING_HEADER_RE.sub("", raw)
            lines = []
            for line in raw.splitlines():
                stripped = line.strip()
                if not stripped or PAGE_NUMBER_RE.match(stripped):
                    continue
                lines.append(stripped)
            pages.append("\n".join(lines))
    finally:
        pdf.close()
    text = "\n".join(pages)
    text = re.sub(r"-\n", "-", text)
    text = RATE_WRAP_RE.sub(r"\1\2", text)
    if document_codes and document_codes[0] not in text:
        text = document_codes[0] + "\n" + text
    return text


def extract_table_code(after_kind: str, collection_number: str) -> str:
    immediate = re.match(r"(\d{2}-\d{2}-\d{3})\b", after_kind.lstrip())
    if immediate:
        return immediate.group(1)

    header = after_kind[:500]
    collapsed = re.sub(r"\s+", "", header)
    collapsed_match = TABLE_CODE_RE.search(collapsed)
    if collapsed_match:
        return collapsed_match.group(0)

    leftover = re.search(r"(?:^|\s)-(\d{2}-\d{3})\b", header)
    if leftover and collection_number:
        return f"{collection_number.zfill(2)}-{leftover.group(1)}"

    fallback = TABLE_CODE_RE.search(header)
    return fallback.group(0) if fallback else ""


def extract_clean_title(header: str) -> str:
    tokens = header.split()
    for i, token in enumerate(tokens):
        if TITLE_WORD_RE.match(token):
            title = normalize_ws(" ".join(tokens[i:]))
            title = re.sub(r"^(?:Таблица\s+ГЭСН[мрп]*\s*)+", "", title)
            return title.strip(" .;")
    return ""


def extend_heading(first_line: str, following: str) -> str:
    title = first_line.strip()
    for line in following.splitlines():
        line = line.strip()
        if not line or MARKER_LINE_RE.match(line):
            break
        should_join = (
            title.endswith("-")
            or title.endswith(",")
            or title.endswith(" И")
            or title.endswith(" ДО")
            or is_mostly_upper(line)
        )
        if should_join:
            glue = "" if title.endswith("-") else " "
            title = title + glue + line
            continue
        break
    return normalize_ws(title)


def parse_composition(block: str) -> list[str]:
    items: list[str] = []
    for raw_line in block.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("Для норм") or line.startswith("Для нормы"):
            continue
        match = COMPOSITION_ITEM_RE.match(line)
        if match:
            items.append(match.group(2).strip())
            continue
        if items and not MARKER_LINE_RE.match(line):
            items[-1] = normalize_ws(items[-1] + " " + line)
    return items


def build_full_name(table_name: str, group: str, name: str) -> str:
    name = normalize_ws(name)
    group = normalize_ws(group).rstrip(":")
    if not name:
        return group or table_name
    if not group:
        if len(name) <= 4 and table_name:
            return f"{table_name}: {name}"
        return name
    if group.lower() in name.lower():
        return name
    return f"{group}: {name}"


def look_like_rate_continuation(line: str) -> bool:
    if not line or line.endswith(":"):
        return False
    if MARKER_LINE_RE.match(line) or COMPOSITION_ITEM_RE.match(line):
        return False
    first = line[0]
    return first.islower() or first.isdigit() or first in "«\"("


def parse_rates(region: str) -> list[tuple[str, str, str]]:
    rates: list[tuple[str, str, str]] = []
    group_buf: list[str] = []
    rates_in_group = 0

    def current_group() -> str:
        return normalize_ws(" ".join(group_buf))

    for raw_line in region.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("Состав работ") or line.startswith("Для норм") or line.startswith("Для нормы"):
            continue
        if COMPOSITION_ITEM_RE.match(line):
            continue
        match = RATE_CODE_RE.match(line)
        if match:
            rates.append((match.group(1), current_group(), match.group(2).strip()))
            rates_in_group += 1
            continue
        if MARKER_LINE_RE.match(line):
            continue
        if rates_in_group and look_like_rate_continuation(line) and line[0].islower():
            code, group, name = rates[-1]
            rates[-1] = (code, group, normalize_ws(name + " " + line))
            continue
        if rates_in_group:
            group_buf = [line]
            rates_in_group = 0
        else:
            group_buf.append(line)
    return rates


def cut_appendices(text: str) -> str:
    first_table = TABLE_RE.search(text)
    if not first_table:
        return text
    appendix = APPENDIX_RE.search(text, first_table.start())
    if appendix:
        return text[: appendix.start()]
    return text


def parse_headings(text: str) -> list[tuple[int, str, str, str]]:
    """Return (offset, kind, number, title) for отдел/раздел/подраздел."""
    headings: list[tuple[int, str, str, str]] = []
    for match in re.finditer(r"^(Отдел|Раздел|Подраздел)\s+", text, re.MULTILINE):
        line_start = match.start()
        line_end = text.find("\n", line_start)
        line = text[line_start:] if line_end < 0 else text[line_start:line_end]
        following = text[line_end + 1 :] if line_end >= 0 else ""
        parsed = OTDEL_RE.match(line) or RAZDEL_RE.match(line) or PODRAZDEL_RE.match(line)
        if not parsed:
            continue
        kind = "otdel" if line.startswith("Отдел") else "razdel" if line.startswith("Раздел") else "podrazdel"
        headings.append((line_start, kind, parsed.group(1), extend_heading(parsed.group(2), following)))
    return headings


def heading_at(headings: list[tuple[int, str, str, str]], offset: int) -> dict[str, str]:
    current = {"department_number": "", "department": "", "section_number": "", "section": "", "subsection_number": "", "subsection": ""}
    for pos, kind, number, title in headings:
        if pos > offset:
            break
        if kind == "otdel":
            current.update(
                department_number=number,
                department=title,
                section_number="",
                section="",
                subsection_number="",
                subsection="",
            )
        elif kind == "razdel":
            current.update(section_number=number, section=title, subsection_number="", subsection="")
        else:
            current.update(subsection_number=number, subsection=title)
    return current


QTY_RE = re.compile(r"\d+(?:,\d+)?")
WORKER_ROW_RE = re.compile(
    r"1-100-(\d{2})\s+Средний разряд работы\s+([\d,]+)\s+чел\.-ч\s*(.*)$"
)
MACHINIST_ROW_RE = re.compile(r"Затраты труда машинистов\s+чел\.-ч\s*(.*)$")
COMMISSIONING_ROW_RE = re.compile(
    r"Затраты труда пусконаладочного персонала[^\n]*?(?:в том числе:)?\s*(.*)$",
    re.IGNORECASE,
)
MACHINE_CODE_RE = re.compile(r"^(91\.\d{2}\.\d{2}-\d{3,4})\b\s*(.*)$")
MACHINE_HOURS_RE = re.compile(r"маш\.-ч\s*(.*)$")
RATE_CODE_TOKEN_RE = re.compile(r"\d{2}-\d{2}-\d{3}-\d{2}")


def parse_quantities(text: str) -> list[float]:
    values: list[float] = []
    for token in QTY_RE.findall(text.replace("\n", " ")):
        values.append(float(token.replace(",", ".")))
    return values


def assign_sequential(columns: list[str], groups: list[tuple[list[float], dict]]) -> dict[str, dict]:
    assigned: dict[str, dict] = {code: {} for code in columns}
    cursor = 0
    for values, extra in groups:
        for value in values:
            if cursor >= len(columns):
                return assigned
            assigned[columns[cursor]] = {"value": value, **extra}
            cursor += 1
    return assigned


def assign_positional(columns: list[str], values: list[float]) -> dict[str, float]:
    result: dict[str, float] = {}
    for code, value in zip(columns, values):
        result[code] = value
    return result


def extract_resource_columns(block: str) -> list[str]:
    columns: list[str] = []
    for raw_line in block.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if re.match(r"^(?:\d\s+)?(?:ЗАТРАТЫ|Затраты труда|МАШИНЫ|МАТЕРИАЛЫ)\b", line):
            break
        if WORKER_ROW_RE.match(line) or line.startswith("91."):
            break
        columns.extend(RATE_CODE_TOKEN_RE.findall(line))
    return columns


def parse_one_resource_table(block: str) -> dict[str, dict]:
    columns = extract_resource_columns(block)
    if not columns:
        return {}

    worker_groups: list[tuple[list[float], dict]] = []
    machinist_values: list[float] = []
    commissioning_values: list[float] = []
    machines: list[tuple[str, str, list[float]]] = []

    pending_machine: tuple[str, str] | None = None
    lines = [line.strip() for line in block.splitlines() if line.strip()]
    skip_materials = False
    for index, line in enumerate(lines):
        if line.startswith("4 МАТЕРИАЛЫ") or line.startswith("МАТЕРИАЛЫ"):
            skip_materials = True
            pending_machine = None
            continue
        if skip_materials:
            continue

        worker_match = WORKER_ROW_RE.match(line)
        if worker_match:
            worker_groups.append(
                (
                    parse_quantities(worker_match.group(3)),
                    {"grade": float(worker_match.group(2).replace(",", "."))},
                )
            )
            pending_machine = None
            continue

        machinist_match = MACHINIST_ROW_RE.search(line)
        if machinist_match:
            machinist_values.extend(parse_quantities(machinist_match.group(1)))
            pending_machine = None
            continue

        if "пусконаладочного персонала" in line.lower():
            tail = COMMISSIONING_ROW_RE.search(line)
            values = parse_quantities(tail.group(1)) if tail else []
            look = index + 1
            while not values and look < len(lines):
                nxt = lines[look]
                look += 1
                if nxt.startswith("2-") or nxt.startswith("3-") or nxt.startswith("Код") or nxt.startswith("91."):
                    break
                if re.search(r"том числе|числе:", nxt, re.I) and not QTY_RE.search(nxt):
                    continue
                values = parse_quantities(nxt)
                if values:
                    break
            commissioning_values.extend(values)
            pending_machine = None
            continue

        hours_match = MACHINE_HOURS_RE.search(line)
        code_match = MACHINE_CODE_RE.match(line)
        if code_match:
            name_part = MACHINE_HOURS_RE.sub("", code_match.group(2)).strip(" ,")
            if hours_match:
                machines.append((code_match.group(1), normalize_ws(name_part), parse_quantities(hours_match.group(1))))
                pending_machine = None
            else:
                pending_machine = (code_match.group(1), name_part)
            continue
        if hours_match and pending_machine:
            code, name = pending_machine
            extra_name = MACHINE_HOURS_RE.sub("", line).strip(" ,")
            full_name = normalize_ws(f"{name} {extra_name}".strip())
            machines.append((code, full_name, parse_quantities(hours_match.group(1))))
            pending_machine = None
            continue
        if pending_machine and not hours_match:
            code, name = pending_machine
            pending_machine = (code, normalize_ws(f"{name} {line}"))

    workers = assign_sequential(columns, worker_groups)
    if len(machinist_values) == len(columns):
        machinists = assign_positional(columns, machinist_values)
    else:
        machinists = {
            code: item["value"]
            for code, item in assign_sequential(columns, [(machinist_values, {})]).items()
            if item
        }
    if len(commissioning_values) == len(columns):
        commissioning = assign_positional(columns, commissioning_values)
    else:
        commissioning = {
            code: item["value"]
            for code, item in assign_sequential(columns, [(commissioning_values, {})]).items()
            if item
        }

    per_rate: dict[str, dict] = {
        code: {
            "workers_hours": workers.get(code, {}).get("value"),
            "workers_grade": workers.get(code, {}).get("grade"),
            "machinists_hours": machinists.get(code),
            "commissioning_hours": commissioning.get(code),
            "machines": [],
            "machines_hours": 0.0,
        }
        for code in columns
    }
    for code, name, values in machines:
        mapped = assign_positional(columns, values) if len(values) == len(columns) else {}
        if not mapped and values:
            mapped = {
                rate_code: item["value"]
                for rate_code, item in assign_sequential(columns, [(values, {})]).items()
                if item
            }
            # Sequential fill for sparse machine rows is unreliable; keep only
            # if the row covers every column or a single leftover column.
            if len(values) not in {1, len(columns)}:
                mapped = {}
        for rate_code, hours in mapped.items():
            if hours is None:
                continue
            per_rate[rate_code]["machines"].append({"code": code, "name": name, "hours": hours})
            per_rate[rate_code]["machines_hours"] += hours
    for payload in per_rate.values():
        if not payload["machines"]:
            payload["machines_hours"] = None
        else:
            payload["machines_hours"] = round(payload["machines_hours"], 4)
    return per_rate


def parse_resources(chunk: str) -> dict[str, dict]:
    resource_start = re.search(r"Код ресурса\b", chunk)
    if not resource_start:
        return {}
    merged: dict[str, dict] = {}
    for block in re.split(r"Код ресурса\b", chunk[resource_start.start() :])[1:]:
        merged.update(parse_one_resource_table(block))
    return merged


@dataclass
class Work:
    code: str
    name: str
    group: str
    full_name: str
    unit: str
    class_id: str
    workers_hours: float | None = None
    workers_grade: float | None = None
    machinists_hours: float | None = None
    commissioning_hours: float | None = None
    machines_hours: float | None = None
    machines: list[dict] = field(default_factory=list)


@dataclass
class Table:
    code: str
    name: str
    kind: str
    unit: str
    composition_text: str
    composition_items: list[str]
    class_id: str
    department_number: str = ""
    department: str = ""
    section_number: str = ""
    section: str = ""
    subsection_number: str = ""
    subsection: str = ""
    works: list[Work] = field(default_factory=list)


@dataclass
class Collection:
    domain_id: str
    domain_name: str
    kind: str
    document_code: str
    number: str
    name: str
    source_file: str
    tables: list[Table] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def parse_collection_meta(text: str, path: Path, domain: dict[str, str]) -> tuple[str, str, str]:
    number = ""
    name = path.stem
    document_code = ""

    file_match = FILENAME_RE.search(path.name)
    if file_match:
        number = file_match.group(2).zfill(2)
        name = file_match.group("name").strip()

    doc_match = DOC_CODE_RE.search(text)
    if doc_match:
        document_code = f"{doc_match.group(1)} {doc_match.group(2)}"

    collection_match = COLLECTION_RE.search(text)
    if collection_match:
        number = collection_match.group(1).zfill(2)
        name = normalize_ws(collection_match.group(2).split("\n")[0])

    if not document_code:
        document_code = f"{domain['kind']} {number}".strip()
    return document_code, number, name


def parse_tables(text: str, kind: str, collection_number: str) -> Iterator[tuple[int, Table]]:
    matches = list(TABLE_RE.finditer(text))
    headings = parse_headings(text)
    for i, match in enumerate(matches):
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        chunk = text[start:end]
        table_kind = f"ГЭСН{match.group(1)}" if match.group(1) else kind
        code = extract_table_code(chunk, collection_number)
        if not code:
            continue

        stop = re.search(r"(Состав работ:|Измеритель:|Код ресурса)", chunk)
        header = chunk[: stop.start()] if stop else chunk[:400]
        title = extract_clean_title(header)
        if not title:
            title = code

        unit_match = re.search(r"Измеритель:\s*(.+)", chunk)
        unit = normalize_ws(unit_match.group(1).split("\n")[0]) if unit_match else ""

        composition_match = re.search(
            r"Состав работ:\s*(.*?)(?:Измеритель:|Код ресурса)",
            chunk,
            re.DOTALL,
        )
        composition_text = normalize_ws(composition_match.group(1)) if composition_match else ""
        composition_items = parse_composition(composition_match.group(1) if composition_match else "")

        rates_region = chunk
        resource_header = re.search(r"\nКод ресурса\b", rates_region)
        if resource_header:
            rates_region = rates_region[: resource_header.start()]
        unit_line = re.search(r"Измеритель:.*", rates_region)
        if unit_line:
            rates_region = rates_region[unit_line.end() :]

        table = Table(
            code=code,
            name=title,
            kind=table_kind,
            unit=unit,
            composition_text=composition_text,
            composition_items=composition_items,
            class_id=f"{table_kind} {code}",
            **heading_at(headings, match.start()),
        )
        seen_codes: set[str] = set()
        resources = parse_resources(chunk)
        for rate_code, group, rate_name in parse_rates(rates_region):
            if rate_code in seen_codes:
                continue
            seen_codes.add(rate_code)
            labor = resources.get(rate_code, {})
            table.works.append(
                Work(
                    code=rate_code,
                    name=rate_name,
                    group=group,
                    full_name=build_full_name(title, group, rate_name),
                    unit=unit,
                    class_id=f"{table_kind} {rate_code}",
                    workers_hours=labor.get("workers_hours"),
                    workers_grade=labor.get("workers_grade"),
                    machinists_hours=labor.get("machinists_hours"),
                    commissioning_hours=labor.get("commissioning_hours"),
                    machines_hours=labor.get("machines_hours"),
                    machines=labor.get("machines") or [],
                )
            )
        yield match.start(), table


def parse_pdf(path: Path, source_root: Path) -> Collection:
    domain = domain_from_folder(path.parent.name)
    text = extract_pdf_text(path)
    text = cut_appendices(text)
    document_code, number, name = parse_collection_meta(text, path, domain)
    try:
        rel = str(path.relative_to(ROOT))
    except ValueError:
        rel = str(path.relative_to(source_root) if source_root in path.parents else path)
    collection = Collection(
        domain_id=domain["id"],
        domain_name=domain["name"],
        kind=domain["kind"],
        document_code=document_code,
        number=number,
        name=name,
        source_file=rel,
    )

    for _, table in parse_tables(text, domain["kind"], number):
        collection.tables.append(table)
    if not collection.tables:
        collection.warnings.append("Не найдено ни одной таблицы ГЭСН")
    return collection


def iter_pdfs(input_dir: Path) -> list[Path]:
    return sorted(p for p in input_dir.rglob("*.pdf") if p.is_file())


def collection_to_dict(collection: Collection) -> dict:
    return {
        "domain_id": collection.domain_id,
        "domain_name": collection.domain_name,
        "kind": collection.kind,
        "document_code": collection.document_code,
        "number": collection.number,
        "name": collection.name,
        "source_file": collection.source_file,
        "warnings": collection.warnings,
        "tables": [
            {
                "class_id": table.class_id,
                "code": table.code,
                "name": table.name,
                "kind": table.kind,
                "unit": table.unit,
                "department_number": table.department_number,
                "department": table.department,
                "section_number": table.section_number,
                "section": table.section,
                "subsection_number": table.subsection_number,
                "subsection": table.subsection,
                "composition": table.composition_text,
                "composition_items": table.composition_items,
                "works": [
                    {
                        "class_id": work.class_id,
                        "code": work.code,
                        "name": work.name,
                        "group": work.group,
                        "full_name": work.full_name,
                        "unit": work.unit,
                        "workers_hours": work.workers_hours,
                        "workers_grade": work.workers_grade,
                        "machinists_hours": work.machinists_hours,
                        "commissioning_hours": work.commissioning_hours,
                        "labor_hours": _labor_hours(work),
                        "machines_hours": work.machines_hours,
                        "machines": work.machines,
                    }
                    for work in table.works
                ],
            }
            for table in collection.tables
        ],
    }


def _labor_hours(work: Work) -> float | None:
    parts = [work.workers_hours, work.machinists_hours, work.commissioning_hours]
    present = [value for value in parts if value is not None]
    if not present:
        return None
    return round(sum(present), 4)


def _format_machines(machines: list[dict]) -> str:
    if not machines:
        return ""
    return "; ".join(f"{item['name']}: {item['hours']}" for item in machines)


def flatten_tables(collections: Iterable[Collection]) -> Iterator[dict[str, str | int]]:
    for collection in collections:
        for table in collection.tables:
            yield {
                "class_id": table.class_id,
                "level": "table",
                "domain_id": collection.domain_id,
                "domain_name": collection.domain_name,
                "kind": collection.kind,
                "document_code": collection.document_code,
                "collection_number": collection.number,
                "collection_name": collection.name,
                "department_number": table.department_number,
                "department": table.department,
                "section_number": table.section_number,
                "section": table.section,
                "subsection_number": table.subsection_number,
                "subsection": table.subsection,
                "table_code": table.code,
                "table_name": table.name,
                "work_code": "",
                "work_name": table.name,
                "full_name": table.name,
                "unit": table.unit,
                "composition": table.composition_text,
                "works_count": len(table.works),
                "source_file": collection.source_file,
            }


def flatten_works(collections: Iterable[Collection]) -> Iterator[dict[str, str | int]]:
    for collection in collections:
        for table in collection.tables:
            for work in table.works:
                yield {
                    "class_id": work.class_id,
                    "level": "work",
                    "domain_id": collection.domain_id,
                    "domain_name": collection.domain_name,
                    "kind": collection.kind,
                    "document_code": collection.document_code,
                    "collection_number": collection.number,
                    "collection_name": collection.name,
                    "department_number": table.department_number,
                    "department": table.department,
                    "section_number": table.section_number,
                    "section": table.section,
                    "subsection_number": table.subsection_number,
                    "subsection": table.subsection,
                    "table_code": table.code,
                    "table_name": table.name,
                    "work_code": work.code,
                    "work_name": work.name,
                    "group": work.group,
                    "full_name": work.full_name,
                    "unit": work.unit,
                    "workers_hours": work.workers_hours,
                    "workers_grade": work.workers_grade,
                    "machinists_hours": work.machinists_hours,
                    "commissioning_hours": work.commissioning_hours,
                    "labor_hours": _labor_hours(work),
                    "machines_hours": work.machines_hours,
                    "machines": _format_machines(work.machines),
                    "source_file": collection.source_file,
                }


WORK_FIELDS = [
    "class_id",
    "level",
    "domain_id",
    "domain_name",
    "kind",
    "document_code",
    "collection_number",
    "collection_name",
    "department_number",
    "department",
    "section_number",
    "section",
    "subsection_number",
    "subsection",
    "table_code",
    "table_name",
    "work_code",
    "work_name",
    "group",
    "full_name",
    "unit",
    "workers_hours",
    "workers_grade",
    "machinists_hours",
    "commissioning_hours",
    "labor_hours",
    "machines_hours",
    "machines",
    "source_file",
]

TABLE_FIELDS = [
    "class_id",
    "level",
    "domain_id",
    "domain_name",
    "kind",
    "document_code",
    "collection_number",
    "collection_name",
    "department_number",
    "department",
    "section_number",
    "section",
    "subsection_number",
    "subsection",
    "table_code",
    "table_name",
    "work_code",
    "work_name",
    "full_name",
    "unit",
    "composition",
    "works_count",
    "source_file",
]


def write_csv(path: Path, rows: list[dict], fieldnames: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, tables: list[dict], works: list[dict], collections: list[Collection]) -> None:
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    ws_works = wb.create_sheet("works")
    ws_works.append(WORK_FIELDS)
    for row in works:
        ws_works.append([row.get(name, "") for name in WORK_FIELDS])

    ws_tables = wb.create_sheet("tables")
    ws_tables.append(TABLE_FIELDS)
    for row in tables:
        ws_tables.append([row.get(name, "") for name in TABLE_FIELDS])

    ws_col = wb.create_sheet("collections")
    ws_col.append(
        ["domain_id", "domain_name", "kind", "document_code", "number", "name", "tables", "works", "source_file"]
    )
    for collection in collections:
        ws_col.append(
            [
                collection.domain_id,
                collection.domain_name,
                collection.kind,
                collection.document_code,
                collection.number,
                collection.name,
                len(collection.tables),
                sum(len(table.works) for table in collection.tables),
                collection.source_file,
            ]
        )
    wb.save(path)


def build_tree(collections: list[Collection]) -> dict:
    domains: dict[str, dict] = {}
    for collection in collections:
        domain = domains.setdefault(
            collection.domain_id,
            {
                "id": collection.domain_id,
                "name": collection.domain_name,
                "kind": collection.kind,
                "collections": [],
            },
        )
        domain["collections"].append(collection_to_dict(collection))
    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": str(DEFAULT_INPUT.relative_to(ROOT)),
        "domains": list(domains.values()),
    }


def parse_all(input_dir: Path, max_files: int | None = None) -> list[Collection]:
    pdfs = iter_pdfs(input_dir)
    if max_files is not None:
        pdfs = pdfs[:max_files]
    collections: list[Collection] = []
    total = len(pdfs)
    for index, path in enumerate(pdfs, 1):
        print(f"[{index}/{total}] {path.name}", file=sys.stderr)
        try:
            collections.append(parse_pdf(path, input_dir))
        except Exception as exc:  # noqa: BLE001 — продолжаем остальные сборники
            print(f"  ошибка: {exc}", file=sys.stderr)
            domain = domain_from_folder(path.parent.name)
            collections.append(
                Collection(
                    domain_id=domain["id"],
                    domain_name=domain["name"],
                    kind=domain["kind"],
                    document_code="",
                    number="",
                    name=path.stem,
                    source_file=str(path.relative_to(ROOT)) if ROOT in path.parents else str(path),
                    warnings=[f"Ошибка парсинга: {exc}"],
                )
            )
    return collections


def save_outputs(collections: list[Collection], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    tree = build_tree(collections)
    tables = list(flatten_tables(collections))
    works = list(flatten_works(collections))

    (output_dir / "classifier.json").write_text(
        json.dumps(tree, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    write_csv(output_dir / "tables.csv", tables, TABLE_FIELDS)
    write_csv(output_dir / "works.csv", works, WORK_FIELDS)
    try:
        write_xlsx(output_dir / "classifier.xlsx", tables, works, collections)
    except ImportError:
        print("openpyxl не установлен — Excel пропущен", file=sys.stderr)

    stats = {
        "generated_at": tree["generated_at"],
        "collections": len(collections),
        "tables": len(tables),
        "works": len(works),
        "collections_without_tables": sum(1 for item in collections if not item.tables),
        "by_domain": {},
    }
    for collection in collections:
        bucket = stats["by_domain"].setdefault(
            collection.domain_id,
            {"name": collection.domain_name, "collections": 0, "tables": 0, "works": 0},
        )
        bucket["collections"] += 1
        bucket["tables"] += len(collection.tables)
        bucket["works"] += sum(len(table.works) for table in collection.tables)
    all_works = [work for collection in collections for table in collection.tables for work in table.works]
    stats["labor"] = {
        "workers_hours": sum(work.workers_hours is not None for work in all_works),
        "machinists_hours": sum(work.machinists_hours is not None for work in all_works),
        "commissioning_hours": sum(work.commissioning_hours is not None for work in all_works),
        "machines_hours": sum(work.machines_hours is not None for work in all_works),
    }
    (output_dir / "stats.json").write_text(
        json.dumps(stats, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(stats, ensure_ascii=False, indent=2))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Собрать классификатор работ из PDF ГЭСН")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Папка с PDF сборников")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Куда сохранить классификатор")
    parser.add_argument("--max-files", type=int, default=None, help="Ограничить число PDF (для отладки)")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.input.exists():
        print(f"Нет папки с документами: {args.input}", file=sys.stderr)
        return 1
    collections = parse_all(args.input, args.max_files)
    save_outputs(collections, args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
